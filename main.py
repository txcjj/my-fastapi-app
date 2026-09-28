"""
餐饮进销存系统 · 后端主程序
FastAPI + SQLAlchemy + PostgreSQL
"""
import os
import io
import csv
import math
import hashlib
import logging
from collections import defaultdict, Counter
from decimal import Decimal, InvalidOperation
from urllib.parse import quote
from contextlib import asynccontextmanager

from fastapi import FastAPI, Depends, HTTPException, Request, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from sqlalchemy import (
    Column, BigInteger, Integer, String, Numeric, Boolean,
    Date, DateTime, JSON, text, or_
)
from sqlalchemy.sql import func
from sqlalchemy.orm import declarative_base, Session
from typing import Optional, List
from datetime import date, datetime, timedelta
from pydantic import BaseModel

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:          # 未安装时，Excel 相关功能给出明确提示，其余接口不受影响
    openpyxl = None

from database import engine, get_db
import schemas

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("inventory-api")

Base = declarative_base()

# ============================================================
#  模型定义（27 张表）
# ============================================================
class Store(Base):
    __tablename__ = "stores"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    name = Column(String(100), nullable=False)
    address = Column(String(255))
    tax_rate = Column(Numeric(5, 4), default=0.01)
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class User(Base):
    __tablename__ = "users"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    auth_uid = Column(String(64), unique=True)
    name = Column(String(50), nullable=False)
    role = Column(String(20), default="staff")
    phone = Column(String(20))
    status = Column(String(20), default="active")
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class Supplier(Base):
    __tablename__ = "suppliers"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    name = Column(String(100), nullable=False)
    contact = Column(String(50))
    phone = Column(String(20))
    address = Column(String(255))
    bank_account = Column(String(50))
    status = Column(String(20), default="active")
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class Warehouse(Base):
    __tablename__ = "warehouses"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    name = Column(String(50), nullable=False)
    type = Column(String(20))
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class Category(Base):
    __tablename__ = "categories"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    name = Column(String(50), nullable=False)
    parent_id = Column(BigInteger)
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class Material(Base):
    __tablename__ = "materials"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    name = Column(String(100), nullable=False)
    category_id = Column(BigInteger)
    unit = Column(String(20), nullable=False, default="kg")
    spec = Column(String(100))
    purchase_unit = Column(String(20))
    purchase_to_base_ratio = Column(Numeric(10, 4), default=1)
    safety_stock = Column(Numeric(10, 4), default=0)
    shelf_life_days = Column(Integer, default=0)
    is_perishable = Column(Boolean, default=False)
    latest_purchase_price = Column(Numeric(10, 4), default=0)
    created_at = Column(DateTime(timezone=True), server_default=func.now())
    updated_at = Column(DateTime(timezone=True), server_default=func.now(), onupdate=func.now())

class Dish(Base):
    __tablename__ = "dishes"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    name = Column(String(100), nullable=False)
    category_id = Column(BigInteger)
    selling_price = Column(Numeric(10, 2), nullable=False, default=0)
    pos_sku = Column(String(50))
    status = Column(String(20), default="active")
    created_at = Column(DateTime(timezone=True), server_default=func.now())
    updated_at = Column(DateTime(timezone=True), server_default=func.now(), onupdate=func.now())

class Bom(Base):
    __tablename__ = "bom"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    dish_id = Column(BigInteger, index=True)
    material_id = Column(BigInteger, index=True)
    quantity = Column(Numeric(10, 4), nullable=False)
    unit = Column(String(20))
    loss_rate = Column(Numeric(5, 4), default=0)
    version = Column(Integer, default=1)
    is_active = Column(Boolean, default=True)
    created_at = Column(DateTime(timezone=True), server_default=func.now())
    updated_at = Column(DateTime(timezone=True), server_default=func.now(), onupdate=func.now())

class PurchaseOrder(Base):
    __tablename__ = "purchase_orders"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    supplier_id = Column(BigInteger, index=True)
    order_no = Column(String(50), unique=True, nullable=False)
    status = Column(String(20), default="pending")
    created_by = Column(BigInteger)
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class PurchaseOrderItem(Base):
    __tablename__ = "purchase_order_items"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    purchase_order_id = Column(BigInteger, index=True)
    material_id = Column(BigInteger, index=True)
    quantity = Column(Numeric(10, 4), nullable=False)
    unit_price = Column(Numeric(10, 4), nullable=False)
    amount = Column(Numeric(12, 2), nullable=False)
    tax_amount = Column(Numeric(12, 2), default=0)
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class StockIn(Base):
    __tablename__ = "stock_in"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    warehouse_id = Column(BigInteger, index=True)
    source_type = Column(String(20), default="采购")
    source_id = Column(BigInteger)
    operator_id = Column(BigInteger)
    status = Column(String(20), default="pending")
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class StockInItem(Base):
    __tablename__ = "stock_in_items"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    stock_in_id = Column(BigInteger, index=True)
    material_id = Column(BigInteger, index=True)
    quantity = Column(Numeric(10, 4), nullable=False)
    unit_cost = Column(Numeric(10, 4), nullable=False)
    batch_no = Column(String(50))
    production_date = Column(Date)
    expire_date = Column(Date)
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class Inventory(Base):
    __tablename__ = "inventory"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    warehouse_id = Column(BigInteger, index=True)
    material_id = Column(BigInteger, index=True)
    quantity = Column(Numeric(12, 4), default=0)
    avg_cost = Column(Numeric(10, 4), default=0)
    last_used_date = Column(Date)
    updated_at = Column(DateTime(timezone=True), server_default=func.now(), onupdate=func.now())

class StockOut(Base):
    __tablename__ = "stock_out"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    warehouse_id = Column(BigInteger, index=True)
    out_type = Column(String(20), default="销售自动扣减")
    source_id = Column(BigInteger)
    operator_id = Column(BigInteger)
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class StockOutItem(Base):
    __tablename__ = "stock_out_items"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    stock_out_id = Column(BigInteger, index=True)
    material_id = Column(BigInteger, index=True)
    quantity = Column(Numeric(10, 4), nullable=False)
    cost_amount = Column(Numeric(12, 2))
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class StockTake(Base):
    __tablename__ = "stock_take"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    warehouse_id = Column(BigInteger, index=True)
    take_date = Column(Date, nullable=False)
    status = Column(String(20), default="pending")
    operator_id = Column(BigInteger)
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class StockTakeItem(Base):
    __tablename__ = "stock_take_items"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    stock_take_id = Column(BigInteger, index=True)
    material_id = Column(BigInteger, index=True)
    book_quantity = Column(Numeric(10, 4), nullable=False)
    actual_quantity = Column(Numeric(10, 4), nullable=False)
    diff_quantity = Column(Numeric(10, 4))
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class PosTemplate(Base):
    __tablename__ = "pos_templates"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    pos_brand = Column(String(50), nullable=False)
    field_mapping = Column(JSON)
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class PosImportBatch(Base):
    __tablename__ = "pos_import_batches"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    file_name = Column(String(255), nullable=False)
    import_date = Column(Date, nullable=False)
    status = Column(String(20), default="pending")
    total_rows = Column(Integer, default=0)
    matched_rows = Column(Integer, default=0)
    unmatched_rows = Column(Integer, default=0)
    # ---- 以下为导入增强字段（启动时自动 ALTER 补列，见 MIGRATIONS） ----
    sale_date_end = Column(Date)                         # 销售日期区间终点（import_date 为起点）
    uploaded_by = Column(String(50))                     # 上传人
    file_hash = Column(String(64))                      # 文件内容 SHA-256，用于防重复
    total_amount = Column(Numeric(14, 2), default=0)     # 销售额合计
    total_quantity = Column(Numeric(14, 2), default=0)   # 销量合计
    stock_deducted = Column(Boolean, default=False)      # 是否已按配方自动扣减库存
    remark = Column(String(255))
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class SalesRecord(Base):
    __tablename__ = "sales_records"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    batch_id = Column(BigInteger, index=True)
    sale_date = Column(Date, nullable=False)
    dish_pos_sku = Column(String(50))
    dish_id = Column(BigInteger, index=True)
    quantity = Column(Numeric(10, 2), nullable=False)
    amount = Column(Numeric(12, 2), nullable=False)
    order_no = Column(String(50))
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class Asset(Base):
    __tablename__ = "assets"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    name = Column(String(100), nullable=False)
    purchase_amount = Column(Numeric(12, 2), nullable=False)
    purchase_date = Column(Date, nullable=False)
    depreciation_method = Column(String(20), default="直线法")
    useful_life_months = Column(Integer, nullable=False)
    residual_rate = Column(Numeric(5, 4), default=0.05)
    monthly_depreciation = Column(Numeric(10, 2))
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class CostDaily(Base):
    __tablename__ = "cost_daily"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    cost_date = Column(Date, nullable=False)
    material_cost = Column(Numeric(12, 2), default=0)
    labor_cost = Column(Numeric(12, 2), default=0)
    rent_cost = Column(Numeric(12, 2), default=0)
    utility_cost = Column(Numeric(12, 2), default=0)
    depreciation_cost = Column(Numeric(12, 2), default=0)
    other_cost = Column(Numeric(12, 2), default=0)
    total_cost = Column(Numeric(12, 2), default=0)

class RevenueDaily(Base):
    __tablename__ = "revenue_daily"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    revenue_date = Column(Date, nullable=False)
    total_revenue = Column(Numeric(12, 2), default=0)
    total_orders = Column(Integer, default=0)
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class ProfitDaily(Base):
    __tablename__ = "profit_daily"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    profit_date = Column(Date, nullable=False)
    revenue = Column(Numeric(12, 2), default=0)
    cost = Column(Numeric(12, 2), default=0)
    gross_profit = Column(Numeric(12, 2), default=0)
    pretax_profit = Column(Numeric(12, 2), default=0)
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class TaxSetting(Base):
    __tablename__ = "tax_settings"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    tax_type = Column(String(50), nullable=False)
    rate = Column(Numeric(6, 4), nullable=False)
    effective_date = Column(Date, nullable=False)
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class TaxMonthly(Base):
    __tablename__ = "tax_monthly"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    period = Column(String(7), nullable=False)
    taxable_income = Column(Numeric(12, 2), default=0)
    tax_payable = Column(Numeric(12, 2), default=0)
    calculation_detail = Column(JSON)
    created_at = Column(DateTime(timezone=True), server_default=func.now())

class TaxYearly(Base):
    __tablename__ = "tax_yearly"
    id = Column(BigInteger, primary_key=True, autoincrement=True)
    store_id = Column(BigInteger, index=True)
    period = Column(String(4), nullable=False)
    taxable_income = Column(Numeric(12, 2), default=0)
    tax_payable = Column(Numeric(12, 2), default=0)
    calculation_detail = Column(JSON)
    created_at = Column(DateTime(timezone=True), server_default=func.now())


# ============================================================
#  补充 Pydantic 模型
# ============================================================
class EnsureMaterialIn(BaseModel):
    store_id: int
    name: str
    unit: Optional[str] = "kg"
    latest_purchase_price: Optional[float] = 0

class PosTemplateIn(BaseModel):
    store_id: int
    pos_brand: str
    field_mapping: Optional[dict] = None

class InventoryIn(BaseModel):
    store_id: Optional[int] = None
    warehouse_id: Optional[int] = None
    material_id: int
    quantity: float = 0
    avg_cost: float = 0
    last_used_date: Optional[date] = None

class PosImportBatchIn(BaseModel):
    store_id: int
    file_name: str
    import_date: date
    status: str = "pending"
    total_rows: int = 0
    matched_rows: int = 0
    unmatched_rows: int = 0
    sale_date_end: Optional[date] = None
    uploaded_by: Optional[str] = None
    file_hash: Optional[str] = None
    total_amount: float = 0
    total_quantity: float = 0
    stock_deducted: bool = False
    remark: Optional[str] = None

class StockTakeLine(BaseModel):
    material_id: int
    actual_quantity: float

class StockTakeSubmitIn(BaseModel):
    store_id: int = 1
    warehouse_id: int = 1
    take_date: date
    operator_id: Optional[int] = None
    items: List[StockTakeLine]

class SalesRecordIn(BaseModel):
    store_id: int
    batch_id: Optional[int] = None
    sale_date: date
    dish_pos_sku: Optional[str] = None
    dish_id: Optional[int] = None
    quantity: float
    amount: float
    order_no: Optional[str] = None

class TaxSettingIn(BaseModel):
    store_id: int
    tax_type: str
    rate: float
    effective_date: date

class TaxMonthlyIn(BaseModel):
    store_id: int
    period: str
    taxable_income: float = 0
    tax_payable: float = 0
    calculation_detail: Optional[dict] = None

class TaxYearlyIn(BaseModel):
    store_id: int
    period: str
    taxable_income: float = 0
    tax_payable: float = 0
    calculation_detail: Optional[dict] = None

# 批量保存配方（新增）
class BomBatchLine(BaseModel):
    material_id: int
    quantity: float
    unit: Optional[str] = None
    loss_rate: float = 0

class BomBatchIn(BaseModel):
    dish_id: int
    items: List[BomBatchLine]


# ============================================================
#  启动/关闭（lifespan 取代已弃用的 on_event）
# ============================================================
# 对已存在的表补列/补索引（create_all 不会修改旧表）。均为幂等语句。
MIGRATIONS = [
    "ALTER TABLE pos_import_batches ADD COLUMN IF NOT EXISTS sale_date_end DATE",
    "ALTER TABLE pos_import_batches ADD COLUMN IF NOT EXISTS uploaded_by VARCHAR(50)",
    "ALTER TABLE pos_import_batches ADD COLUMN IF NOT EXISTS file_hash VARCHAR(64)",
    "ALTER TABLE pos_import_batches ADD COLUMN IF NOT EXISTS total_amount NUMERIC(14,2) DEFAULT 0",
    "ALTER TABLE pos_import_batches ADD COLUMN IF NOT EXISTS total_quantity NUMERIC(14,2) DEFAULT 0",
    "ALTER TABLE pos_import_batches ADD COLUMN IF NOT EXISTS stock_deducted BOOLEAN DEFAULT FALSE",
    "ALTER TABLE pos_import_batches ADD COLUMN IF NOT EXISTS remark VARCHAR(255)",
    "CREATE INDEX IF NOT EXISTS ix_pos_import_batches_file_hash ON pos_import_batches (store_id, file_hash)",
    "CREATE INDEX IF NOT EXISTS ix_sales_records_store_date ON sales_records (store_id, sale_date)",
]


@asynccontextmanager
async def lifespan(app: FastAPI):
    # ---- 启动 ----
    try:
        Base.metadata.create_all(bind=engine)
        logger.info("✅ 数据库表结构已就绪")
    except Exception as e:
        logger.error(f"❌ 建表失败：{e}")
    for sql in MIGRATIONS:
        try:
            with engine.begin() as conn:
                conn.execute(text(sql))
        except Exception as e:
            logger.warning(f"迁移语句跳过：{sql[:60]}... -> {e}")
    yield
    # ---- 关闭（暂无资源需要释放） ----


# ============================================================
#  FastAPI 应用
# ============================================================
app = FastAPI(
    title="餐饮进销存系统 API",
    description="27 张业务表的 RESTful 接口",
    version="1.0.0",
    docs_url="/docs",
    redoc_url="/redoc",
    lifespan=lifespan,
)

# CORS：默认允许所有来源；生产环境可通过环境变量 CORS_ORIGINS（逗号分隔）收敛
_cors_env = os.environ.get("CORS_ORIGINS", "").strip()
_cors_origins = [o.strip() for o in _cors_env.split(",") if o.strip()] if _cors_env else ["*"]

app.add_middleware(
    CORSMiddleware,
    allow_origins=_cors_origins,
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ============================================================
#  通用 CRUD 注册器
# ============================================================
def _coerce(col, v: str):
    """把查询字符串按列类型转换，便于 ?purchase_order_id=3 这类过滤"""
    t = col.type
    try:
        if isinstance(t, (Integer, BigInteger)):
            return int(v)
        if isinstance(t, Boolean):
            return str(v).lower() in ("1", "true", "yes")
        if isinstance(t, Date) and not isinstance(t, DateTime):
            return date.fromisoformat(v)
        if isinstance(t, Numeric):
            return Decimal(v)
    except (ValueError, InvalidOperation):
        raise HTTPException(status_code=400, detail=f"参数 {col.name} 格式不正确")
    return v


def register_crud(model, schema_in, prefix: str, tag: str, search_field: str = None, skip_ops: tuple = ()):
    """通用 CRUD。skip_ops 可传 ("update", "delete") 等，跳过自动生成、改由下方自定义接口实现。
    列表接口支持：keyword 模糊搜索、任意列等值过滤（?列名=值）、skip/limit 分页，默认按 id 倒序。"""
    @app.get(f"/api{prefix}", tags=[tag], summary=f"获取 {tag} 列表")
    def list_items(
        request: Request,
        keyword: Optional[str] = None,
        skip: int = 0,
        limit: int = 1000,
        db: Session = Depends(get_db),
    ):
        q = db.query(model)
        if keyword and search_field and hasattr(model, search_field):
            q = q.filter(getattr(model, search_field).ilike(f"%{keyword}%"))
        for k, v in request.query_params.items():
            if k in ("keyword", "skip", "limit"):
                continue
            col = model.__table__.columns.get(k)
            if col is not None:
                q = q.filter(col == _coerce(col, v))
        # 上限提高到 20000，避免盘点明细等大表被截断
        return q.order_by(model.id.desc()).offset(skip).limit(min(limit, 20000)).all()

    @app.get(f"/api{prefix}/{{item_id}}", tags=[tag], summary=f"获取 {tag} 详情")
    def get_item(item_id: int, db: Session = Depends(get_db)):
        obj = db.query(model).filter(model.id == item_id).first()
        if not obj:
            raise HTTPException(status_code=404, detail=f"{tag} 记录不存在")
        return obj

    @app.post(f"/api{prefix}", tags=[tag], summary=f"新增 {tag}", status_code=201)
    def create_item(payload: schema_in, db: Session = Depends(get_db)):
        obj = model(**payload.model_dump(exclude_unset=True))
        db.add(obj)
        db.commit()
        db.refresh(obj)
        return obj

    if "update" not in skip_ops:
        @app.put(f"/api{prefix}/{{item_id}}", tags=[tag], summary=f"更新 {tag}")
        def update_item(item_id: int, payload: schema_in, db: Session = Depends(get_db)):
            obj = db.query(model).filter(model.id == item_id).first()
            if not obj:
                raise HTTPException(status_code=404, detail=f"{tag} 记录不存在")
            for k, v in payload.model_dump(exclude_unset=True).items():
                setattr(obj, k, v)
            db.commit()
            db.refresh(obj)
            return obj

    if "delete" not in skip_ops:
        @app.delete(f"/api{prefix}/{{item_id}}", tags=[tag], summary=f"删除 {tag}")
        def delete_item(item_id: int, db: Session = Depends(get_db)):
            obj = db.query(model).filter(model.id == item_id).first()
            if not obj:
                raise HTTPException(status_code=404, detail=f"{tag} 记录不存在")
            db.delete(obj)
            db.commit()
            return {"ok": True, "deleted_id": item_id}


# 注册所有表
register_crud(Store,             schemas.StoreIn,             "/stores",              "门店",        "name")
register_crud(User,              schemas.UserIn,              "/users",               "用户",        "name")
register_crud(Supplier,          schemas.SupplierIn,          "/suppliers",           "供应商",      "name")
register_crud(Warehouse,         schemas.WarehouseIn,         "/warehouses",          "仓库",        "name")
register_crud(Category,          schemas.CategoryIn,          "/categories",          "分类",        "name")
register_crud(Material,          schemas.MaterialIn,          "/materials",           "原材料",      "name")
register_crud(Dish,              schemas.DishIn,              "/dishes",              "菜品",        "name")
register_crud(Bom,               schemas.BomItemIn,           "/bom",                 "配方")
register_crud(PurchaseOrder,     schemas.PurchaseOrderIn,     "/purchase-orders",     "采购单",      "order_no", skip_ops=("update", "delete"))
register_crud(StockIn,           schemas.StockInIn,           "/stock-in",            "入库单",      None, skip_ops=("update", "delete"))
register_crud(PurchaseOrderItem, schemas.PurchaseItemIn,      "/purchase-order-items","采购单明细")
register_crud(StockInItem,       schemas.StockInItemIn,       "/stock-in-items",      "入库明细",    "batch_no")
register_crud(Inventory,         InventoryIn,                 "/inventory",           "库存")
register_crud(StockOut,          schemas.StockOutIn,          "/stock-out",           "出库单")
register_crud(StockOutItem,      schemas.StockOutItemIn,      "/stock-out-items",     "出库明细")
register_crud(StockTake,         schemas.StockTakeIn,         "/stock-take",          "盘点单")
register_crud(StockTakeItem,     schemas.StockTakeItemIn,     "/stock-take-items",    "盘点明细")
register_crud(PosTemplate,       PosTemplateIn,               "/pos-templates",       "POS 模板",    "pos_brand")
register_crud(PosImportBatch,    PosImportBatchIn,            "/pos-import-batches",  "POS 导入批次", "file_name", skip_ops=("delete",))
register_crud(SalesRecord,       SalesRecordIn,               "/sales-records",       "销售流水",    "order_no")
register_crud(Asset,             schemas.AssetIn,             "/assets",              "固定资产",    "name")
register_crud(CostDaily,         schemas.CostDailyIn,         "/cost-daily",          "每日成本")
register_crud(RevenueDaily,      schemas.RevenueDailyIn,      "/revenue-daily",       "每日营收")
register_crud(ProfitDaily,       schemas.ProfitDailyIn,       "/profit-daily",        "每日利润")
register_crud(TaxSetting,        TaxSettingIn,                "/tax-settings",        "税务参数",    "tax_type")
register_crud(TaxMonthly,        TaxMonthlyIn,                "/tax-monthly",         "月度税务",    "period")
register_crud(TaxYearly,         TaxYearlyIn,                 "/tax-yearly",          "年度税务",    "period")


# ============================================================
#  自定义接口
#  注意：Starlette 按注册顺序匹配路由，先注册者优先，所以"覆盖"通用接口
#  不能靠写在后面，而是在 register_crud 里用 skip_ops 跳过对应的自动接口。
# ============================================================

# ---------- 1. 采购单：审批后禁止删改 ----------
@app.delete("/api/purchase-orders/{item_id}", tags=["采购单"])
def delete_purchase_order(item_id: int, db: Session = Depends(get_db)):
    po = db.query(PurchaseOrder).filter(PurchaseOrder.id == item_id).first()
    if not po:
        raise HTTPException(status_code=404, detail="采购单不存在")
    if po.status == 'approved':
        raise HTTPException(status_code=400, detail="已审批的采购单不可删除")
    db.query(PurchaseOrderItem).filter(PurchaseOrderItem.purchase_order_id == item_id).delete()
    db.delete(po)
    db.commit()
    return {"ok": True}


@app.put("/api/purchase-orders/{item_id}", tags=["采购单"])
def update_purchase_order(item_id: int, payload: schemas.PurchaseOrderIn, db: Session = Depends(get_db)):
    po = db.query(PurchaseOrder).filter(PurchaseOrder.id == item_id).first()
    if not po:
        raise HTTPException(status_code=404, detail="采购单不存在")
    if po.status == 'approved':
        raise HTTPException(status_code=400, detail="已审批的采购单不可修改")
    for k, v in payload.model_dump(exclude_unset=True).items():
        setattr(po, k, v)
    db.commit()
    db.refresh(po)
    return po


# ---------- 2. 采购单能否入库 ----------
@app.get("/api/purchase-orders/{po_id}/can-stock-in", tags=["入库单"])
def can_stock_in(po_id: int, db: Session = Depends(get_db)):
    po = db.query(PurchaseOrder).filter(PurchaseOrder.id == po_id).first()
    if not po:
        return {"ok": False, "reason": "采购单不存在"}
    if po.status != 'approved':
        return {"ok": False, "reason": "采购单尚未审批"}
    existing = db.query(StockIn).filter(
        StockIn.source_id == po_id,
        StockIn.source_type == '采购',
        StockIn.status == 'approved'
    ).first()
    if existing:
        return {"ok": False, "reason": f"已于入库单 #{existing.id} 入库"}
    pending = db.query(StockIn).filter(
        StockIn.source_id == po_id,
        StockIn.source_type == '采购',
        StockIn.status == 'pending'
    ).first()
    if pending:
        return {"ok": False, "reason": f"已有待审批入库单 #{pending.id}"}
    return {"ok": True, "reason": ""}


# ---------- 3. 入库单：审批 + 写库存（一步到位） ----------
@app.post("/api/stock-in/{item_id}/approve", tags=["入库单"])
def approve_stock_in(item_id: int, db: Session = Depends(get_db)):
    si = db.query(StockIn).filter(StockIn.id == item_id).first()
    if not si:
        raise HTTPException(status_code=404, detail="入库单不存在")
    if si.status == 'approved':
        raise HTTPException(status_code=400, detail="该入库单已审批，不可重复操作")
    if si.status == 'rejected':
        raise HTTPException(status_code=400, detail="该入库单已被驳回，请重新提交")

    # 检查采购单是否已入库
    existing = db.query(StockIn).filter(
        StockIn.source_id == si.source_id,
        StockIn.source_type == '采购',
        StockIn.status == 'approved',
        StockIn.id != si.id
    ).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"该采购单已完成入库（入库单 #{existing.id}）")

    # 改状态
    si.status = 'approved'

    # 写库存
    items = db.query(StockInItem).filter(StockInItem.stock_in_id == item_id).all()
    for it in items:
        inv = db.query(Inventory).filter(
            Inventory.store_id == si.store_id, Inventory.material_id == it.material_id
        ).first()
        if inv:
            # 负库存保护：加权平均时账面负数按 0 参与，避免成本被拉歪；
            # 但 quantity 仍按实际（含负数）累加，保留真实账面。
            old_qty_actual = float(inv.quantity or 0)
            old_qty = max(old_qty_actual, 0)
            old_cost = float(inv.avg_cost or 0)
            add_qty = float(it.quantity or 0)
            unit_cost = float(it.unit_cost or 0)
            new_qty = old_qty + add_qty
            new_cost = (old_qty * old_cost + add_qty * unit_cost) / new_qty if new_qty > 0 else unit_cost
            inv.quantity = old_qty_actual + add_qty
            inv.avg_cost = new_cost
        else:
            db.add(Inventory(
                store_id=si.store_id,
                warehouse_id=si.warehouse_id,
                material_id=it.material_id,
                quantity=it.quantity,
                avg_cost=it.unit_cost,
            ))
    db.commit()
    return {"ok": True, "stock_in_id": item_id, "items": len(items)}


# ---------- 4. 入库单：审批后禁止删改 ----------
@app.delete("/api/stock-in/{item_id}", tags=["入库单"])
def delete_stock_in(item_id: int, db: Session = Depends(get_db)):
    si = db.query(StockIn).filter(StockIn.id == item_id).first()
    if not si:
        raise HTTPException(status_code=404, detail="入库单不存在")
    if si.status == 'approved':
        raise HTTPException(status_code=400, detail="已入库的入库单不可删除")
    db.query(StockInItem).filter(StockInItem.stock_in_id == item_id).delete()
    db.delete(si)
    db.commit()
    return {"ok": True}


@app.put("/api/stock-in/{item_id}", tags=["入库单"])
def update_stock_in(item_id: int, payload: schemas.StockInIn, db: Session = Depends(get_db)):
    si = db.query(StockIn).filter(StockIn.id == item_id).first()
    if not si:
        raise HTTPException(status_code=404, detail="入库单不存在")
    if si.status == 'approved':
        raise HTTPException(status_code=400, detail="已入库的入库单不可修改")
    for k, v in payload.model_dump(exclude_unset=True).items():
        setattr(si, k, v)
    db.commit()
    db.refresh(si)
    return si


# ---------- 5. 按名称确保原材料存在（采购时自动建档） ----------
@app.post("/api/materials-ensure", tags=["原材料"])
def ensure_material(payload: EnsureMaterialIn, db: Session = Depends(get_db)):
    name = payload.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="名称不能为空")

    exist = db.query(Material).filter(
        Material.store_id == payload.store_id,
        Material.name == name
    ).first()

    if exist:
        if payload.latest_purchase_price and payload.latest_purchase_price > 0:
            exist.latest_purchase_price = payload.latest_purchase_price
            db.commit()
            db.refresh(exist)
        return exist

    new_mat = Material(
        store_id=payload.store_id,
        name=name,
        unit=payload.unit or "kg",
        latest_purchase_price=payload.latest_purchase_price or 0,
    )
    db.add(new_mat)
    db.commit()
    db.refresh(new_mat)
    return new_mat


# ---------- 5.1 配方：整单批量保存（事务内软删旧 + 插新，避免中途失败丢配方） ----------
@app.put("/api/bom/batch", tags=["配方"], summary="按菜品批量保存配方（事务）")
def save_bom_batch(payload: BomBatchIn, db: Session = Depends(get_db)):
    dish = db.query(Dish).filter(Dish.id == payload.dish_id).first()
    if not dish:
        raise HTTPException(status_code=404, detail="菜品不存在")

    items = [it for it in payload.items if it.material_id and it.quantity and it.quantity > 0]
    if not items:
        raise HTTPException(status_code=400, detail="请至少填写一条有效配方")

    try:
        # 软删该菜品现有生效配方
        old = db.query(Bom).filter(
            Bom.dish_id == payload.dish_id,
            _bom_filter()
        ).all()
        for b in old:
            b.is_active = False
        db.flush()

        # 插入新配方
        created = []
        for it in items:
            b = Bom(
                dish_id=payload.dish_id,
                material_id=it.material_id,
                quantity=it.quantity,
                unit=it.unit,
                loss_rate=it.loss_rate or 0,
                version=1,
                is_active=True,
            )
            db.add(b)
            created.append(b)
        db.commit()
    except Exception as e:
        db.rollback()
        logger.exception("批量保存配方失败")
        raise HTTPException(status_code=500, detail=f"保存配方失败，已整体回滚：{e}")

    return {"ok": True, "dish_id": payload.dish_id, "deactivated": len(old), "created": len(created)}


# ============================================================
#  6. 销售数据导入（模板 → 选择文件 → 校验预览 → 防重复 → 入库 → 自动记录）
# ============================================================
IMPORT_COLUMNS = [
    # (列名, 必填, 填写规则, 示例)
    ("销售日期", True,  "格式 YYYY-MM-DD，不能晚于今天",              "2026-09-24"),
    ("POS编码",  True,  "须与「菜品与配方」里的 POS 编码一致",        "FR001"),
    ("菜品名称", False, "仅供核对，不参与匹配",                      "扬州炒饭"),
    ("数量",     True,  "大于 0 的数字",                            "2"),
    ("金额",     True,  "实收金额（元），不小于 0",                  "36.00"),
    ("订单号",   False, "POS 小票号，填写后可精确防重复",             "T20260924001"),
]
IMPORT_HEADERS = [c[0] for c in IMPORT_COLUMNS]
IMPORT_REQUIRED = [c[0] for c in IMPORT_COLUMNS if c[1]]
SAMPLE_ROWS = [
    ["2026-09-24", "FR001", "扬州炒饭（示例，请删除）", 2, 36.00, "T20260924001"],
    ["2026-09-24", "FR002", "蛋炒饭（示例，请删除）",   1, 15.00, "T20260924002"],
    ["2026-09-24", "DR001", "酸梅汤（示例，请删除）",   3, 24.00, "T20260924002"],
]
MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 20000


class ImportFormatError(Exception):
    pass


def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _cell_str(v).split(" ")[0].split("T")[0]
    if not s:
        return None
    for f in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d", "%Y年%m月%d日"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            continue
    return None


def _to_decimal(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float, Decimal)):
        d = Decimal(str(v))
    else:
        s = str(v).strip().replace(",", "").replace("¥", "").replace("￥", "")
        if not s:
            return None
        try:
            d = Decimal(s)
        except InvalidOperation:
            return None
    return d if d.is_finite() else None


def _read_table(filename: str, content: bytes) -> list:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        if openpyxl is None:
            raise ImportFormatError("服务器未安装 openpyxl，暂不能解析 Excel（pip install openpyxl），或改用 CSV 文件")
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
            wb.close()
            return rows
        except Exception:
            raise ImportFormatError("无法读取该 Excel 文件，文件可能已损坏或不是标准 .xlsx 格式")
    if name.endswith(".csv"):
        for enc in ("utf-8-sig", "gb18030"):
            try:
                return list(csv.reader(io.StringIO(content.decode(enc))))
            except UnicodeDecodeError:
                continue
        raise ImportFormatError("CSV 编码无法识别，请另存为 UTF-8 或 GBK 编码")
    raise ImportFormatError("仅支持 .xlsx 或 .csv 文件（旧版 .xls 请先另存为 .xlsx）")


def _parse_sales_file(filename: str, content: bytes) -> dict:
    """解析并逐行校验。不访问数据库。"""
    res = {"ok": False, "file_errors": [], "row_errors": [], "error_total": 0, "rows": []}
    if not content:
        res["file_errors"].append("文件是空的")
    elif len(content) > MAX_IMPORT_BYTES:
        res["file_errors"].append(f"文件超过 {MAX_IMPORT_BYTES // 1024 // 1024}MB 上限")
    if res["file_errors"]:
        res["error_total"] = len(res["file_errors"])
        return res
    try:
        table = _read_table(filename, content)
    except ImportFormatError as e:
        res["file_errors"].append(str(e))
        res["error_total"] = 1
        return res

    hdr_idx = next((i for i, r in enumerate(table) if any(_cell_str(c) for c in r)), None)
    if hdr_idx is None:
        res["file_errors"].append("文件中没有任何内容")
        res["error_total"] = 1
        return res
    header = [_cell_str(c).replace("*", "").strip() for c in table[hdr_idx]]
    pos = {h: i for i, h in enumerate(header) if h}
    missing = [h for h in IMPORT_REQUIRED if h not in pos]
    if missing:
        res["file_errors"].append(f"缺少必填列：{'、'.join(missing)}（表头必须与模板一致：{'、'.join(IMPORT_HEADERS)}）")
        res["error_total"] = 1
        return res

    def get(row, label):
        i = pos.get(label)
        return row[i] if i is not None and i < len(row) else None

    today = date.today()
    errors, err_count, rows = [], 0, []

    def add_err(row_no, col, msg):
        nonlocal err_count
        err_count += 1
        if len(errors) < 100:
            errors.append({"row": row_no, "column": col, "message": msg})

    for idx in range(hdr_idx + 1, len(table)):
        raw = table[idx]
        if not any(_cell_str(c) for c in raw):
            continue
        row_no = idx + 1
        ok = True
        d = _to_date(get(raw, "销售日期"))
        if d is None:
            add_err(row_no, "销售日期", "日期格式不正确，应为 YYYY-MM-DD"); ok = False
        elif d > today:
            add_err(row_no, "销售日期", "销售日期不能晚于今天"); ok = False
        elif d.year < 2000:
            add_err(row_no, "销售日期", "销售日期不合理"); ok = False
        sku = _cell_str(get(raw, "POS编码"))
        if not sku:
            add_err(row_no, "POS编码", "不能为空"); ok = False
        elif len(sku) > 50:
            add_err(row_no, "POS编码", "长度不能超过 50"); ok = False
        name = _cell_str(get(raw, "菜品名称"))
        if "示例" in name:
            add_err(row_no, "菜品名称", "这是模板里的示例数据，请删除后再上传"); ok = False
        qty = _to_decimal(get(raw, "数量"))
        if qty is None or qty <= 0 or qty >= Decimal("100000000"):
            add_err(row_no, "数量", "必须是大于 0 的数字"); ok = False
        amt = _to_decimal(get(raw, "金额"))
        if amt is None or amt < 0 or amt >= Decimal("10000000000"):
            add_err(row_no, "金额", "必须是不小于 0 的数字"); ok = False
        order_no = _cell_str(get(raw, "订单号"))
        if len(order_no) > 50:
            add_err(row_no, "订单号", "长度不能超过 50"); ok = False
        if ok:
            rows.append({"row_no": row_no, "sale_date": d, "pos_sku": sku, "dish_name": name,
                         "quantity": qty.quantize(Decimal("0.01")), "amount": amt.quantize(Decimal("0.01")),
                         "order_no": order_no})
        if len(rows) + err_count > MAX_IMPORT_ROWS:
            res["file_errors"].append(f"数据行数超过 {MAX_IMPORT_ROWS} 行上限，请拆分文件")
            break
    if not rows and not err_count and not res["file_errors"]:
        res["file_errors"].append("文件中没有数据行")
    res["row_errors"] = errors
    res["error_total"] = len(res["file_errors"]) + err_count
    res["rows"] = rows
    res["ok"] = res["error_total"] == 0
    return res


def _format_error_response(parsed: dict) -> dict:
    return {"ok": False, "code": "format_error", "message": "数据格式错误，请重新编辑",
            "file_errors": parsed["file_errors"], "row_errors": parsed["row_errors"],
            "error_total": parsed["error_total"]}


def _dish_sku_map(db: Session, store_id: int) -> dict:
    dishes = db.query(Dish).filter(Dish.store_id == store_id).all()
    m = {}
    for d in sorted(dishes, key=lambda x: (x.status != "active", x.id)):     # 在售菜品优先
        key = (d.pos_sku or "").strip().upper()
        if key:
            m.setdefault(key, d)
    return m


def _sig(sale_date, sku, qty, amt, order_no):
    q = Decimal(str(qty)).quantize(Decimal("0.01"))
    a = Decimal(str(amt)).quantize(Decimal("0.01"))
    return (sale_date, (sku or "").strip().upper(), q, a, (order_no or "").strip())


def _check_duplicate(db: Session, store_id: int, rows: list, file_hash: str):
    """返回 None（无重复）或 {"level": "full"|"partial", ...}"""
    hit = db.query(PosImportBatch).filter(
        PosImportBatch.store_id == store_id, PosImportBatch.file_hash == file_hash).first()
    if hit:
        return {"level": "full", "batch": hit, "dup_rows": len(rows)}
    dmin = min(r["sale_date"] for r in rows)
    dmax = max(r["sale_date"] for r in rows)
    existing = db.query(SalesRecord.sale_date, SalesRecord.dish_pos_sku, SalesRecord.quantity,
                        SalesRecord.amount, SalesRecord.order_no).filter(
        SalesRecord.store_id == store_id, SalesRecord.sale_date >= dmin, SalesRecord.sale_date <= dmax).all()
    if not existing:
        return None
    exist_cnt = Counter(_sig(*r) for r in existing)
    file_cnt = Counter(_sig(r["sale_date"], r["pos_sku"], r["quantity"], r["amount"], r["order_no"]) for r in rows)
    if all(exist_cnt[k] >= v for k, v in file_cnt.items()):
        return {"level": "full", "batch": None, "dup_rows": len(rows)}
    # 部分重复：只有带订单号的行才可判定为"同一笔订单"，避免无订单号时误伤
    dup_rows = sum(min(exist_cnt[k], v) for k, v in file_cnt.items() if k[4])
    if dup_rows > 0:
        return {"level": "partial", "batch": None, "dup_rows": dup_rows}
    return None


def _refresh_revenue_daily(db: Session, store_id: int, dates):
    for d in dates:
        rev, cnt, orders = db.query(
            func.coalesce(func.sum(SalesRecord.amount), 0), func.count(SalesRecord.id),
            func.count(func.distinct(SalesRecord.order_no)),
        ).filter(SalesRecord.store_id == store_id, SalesRecord.sale_date == d).one()
        row = db.query(RevenueDaily).filter(RevenueDaily.store_id == store_id, RevenueDaily.revenue_date == d).first()
        order_cnt = int(orders) if orders else int(cnt)
        if row:
            row.total_revenue = rev
            row.total_orders = order_cnt
        elif cnt:
            db.add(RevenueDaily(store_id=store_id, revenue_date=d, total_revenue=rev, total_orders=order_cnt))


def _bom_filter():
    return or_(Bom.is_active.is_(True), Bom.is_active.is_(None))


def _dish_bom(db: Session, store_id: int) -> dict:
    """dish_id -> [(material_id, 单份用量, 损耗率)]"""
    ids = [x[0] for x in db.query(Dish.id).filter(Dish.store_id == store_id).all()]
    out = defaultdict(list)
    if ids:
        for b in db.query(Bom).filter(_bom_filter(), Bom.dish_id.in_(ids)).all():
            out[b.dish_id].append((b.material_id, float(b.quantity or 0), float(b.loss_rate or 0)))
    return out


def _material_cost_map(db: Session, store_id: int):
    mats = {m.id: m for m in db.query(Material).filter(Material.store_id == store_id).all()}
    inv_cost = {}
    for i in db.query(Inventory).filter(Inventory.store_id == store_id).all():
        c = float(i.avg_cost or 0)
        if c > 0:
            inv_cost.setdefault(i.material_id, c)
    cost = {mid: inv_cost.get(mid) or float(m.latest_purchase_price or 0) for mid, m in mats.items()}
    return mats, cost


def _deduct_stock(db: Session, store_id: int, batch_id: int, sold: list) -> dict:
    """按配方（含损耗率）扣减库存并生成出库单。sold: [(dish_id, quantity, sale_date)]"""
    bom = _dish_bom(db, store_id)
    need = defaultdict(float)
    no_bom = set()
    last_date = None
    for dish_id, qty, sd in sold:
        if dish_id is None:
            continue
        lines = bom.get(dish_id)
        if not lines:
            no_bom.add(dish_id)
            continue
        for mid, q, loss in lines:
            need[mid] += float(qty) * q * (1 + loss)
        last_date = max(last_date, sd) if last_date else sd
    dish_names = {d.id: d.name for d in db.query(Dish).filter(Dish.id.in_(list(no_bom))).all()} if no_bom else {}
    info = {"materials": 0, "cost_amount": 0.0, "shortage": [], "dishes_without_bom": sorted(dish_names.values())}
    if not need:
        return info
    mats, cost = _material_cost_map(db, store_id)
    wh = db.query(Warehouse).filter(Warehouse.store_id == store_id).order_by(Warehouse.id).first()
    default_wh = wh.id if wh else 1
    so = StockOut(store_id=store_id, warehouse_id=default_wh, out_type="销售自动扣减", source_id=batch_id)
    db.add(so)
    db.flush()
    for mid, qty in need.items():
        inv = db.query(Inventory).filter(Inventory.store_id == store_id, Inventory.material_id == mid) \
            .order_by(Inventory.id).first()
        unit_cost = float(inv.avg_cost or 0) if inv and float(inv.avg_cost or 0) > 0 else cost.get(mid, 0.0)
        if inv:
            inv.quantity = float(inv.quantity or 0) - qty
            inv.last_used_date = last_date
        else:
            inv = Inventory(store_id=store_id, warehouse_id=default_wh, material_id=mid,
                            quantity=-qty, avg_cost=unit_cost, last_used_date=last_date)
            db.add(inv)
        amount = round(qty * unit_cost, 2)
        db.add(StockOutItem(stock_out_id=so.id, material_id=mid, quantity=round(qty, 4), cost_amount=amount))
        info["materials"] += 1
        info["cost_amount"] += amount
        if float(inv.quantity or 0) < 0:
            m = mats.get(mid)
            info["shortage"].append({"material": m.name if m else str(mid), "stock": round(float(inv.quantity), 3)})
    info["cost_amount"] = round(info["cost_amount"], 2)
    return info


def _rollback_batch(db: Session, batch: PosImportBatch):
    """撤销一个导入批次：还原库存、删除出库单与销售流水、刷新日营收"""
    store_id = batch.store_id
    dates = {d for (d,) in db.query(SalesRecord.sale_date).filter(SalesRecord.batch_id == batch.id).distinct()}
    outs = db.query(StockOut).filter(StockOut.store_id == batch.store_id, StockOut.out_type == "销售自动扣减",
                                     StockOut.source_id == batch.id).all()
    for so in outs:
        for it in db.query(StockOutItem).filter(StockOutItem.stock_out_id == so.id).all():
            inv = db.query(Inventory).filter(Inventory.store_id == batch.store_id,
                                             Inventory.material_id == it.material_id).order_by(Inventory.id).first()
            if inv:
                inv.quantity = float(inv.quantity or 0) + float(it.quantity or 0)
            db.delete(it)
        db.delete(so)
    db.query(SalesRecord).filter(SalesRecord.batch_id == batch.id).delete(synchronize_session=False)
    db.delete(batch)
    db.flush()
    _refresh_revenue_daily(db, store_id, dates)


# ---------- 6.1 模板信息 / 模板下载 ----------
@app.get("/api/sales-import/template-info", tags=["销售导入"], summary="模板预览信息")
def sales_template_info():
    return {
        "columns": [{"name": n, "required": r, "rule": rule, "example": ex} for n, r, rule, ex in IMPORT_COLUMNS],
        "sample_rows": SAMPLE_ROWS,
        "max_rows": MAX_IMPORT_ROWS,
        "max_mb": MAX_IMPORT_BYTES // 1024 // 1024,
    }


@app.get("/api/sales-import/template", tags=["销售导入"], summary="下载导入模板（xlsx / csv）")
def sales_template_download(fmt: str = "xlsx"):
    fmt = fmt.lower()
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(IMPORT_HEADERS)
        w.writerows(SAMPLE_ROWS)
        data, mime, fname = buf.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8", "销售数据导入模板.csv"
    else:
        if openpyxl is None:
            raise HTTPException(status_code=500, detail="服务器未安装 openpyxl，请 pip install openpyxl 或下载 CSV 模板")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "销售数据"
        ws.append(IMPORT_HEADERS)
        for r in SAMPLE_ROWS:
            ws.append(r)
        head_fill = PatternFill("solid", fgColor="2F6F4E")
        for i, (name, req, _, _) in enumerate(IMPORT_COLUMNS, start=1):
            c = ws.cell(row=1, column=i)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(i)].width = 18 if i != 3 else 26
        for r in range(2, 2001):                      # 日期 / POS编码 / 订单号 设为文本，防止 Excel 吞掉前导 0
            for col in (1, 2, 6):
                ws.cell(row=r, column=col).number_format = "@"
        ws.freeze_panes = "A2"
        ws2 = wb.create_sheet("填写说明")
        ws2.append(["列名", "是否必填", "填写规则", "示例"])
        for n, req, rule, ex in IMPORT_COLUMNS:
            ws2.append([n, "必填" if req else "选填", rule, ex])
        ws2.append([])
        ws2.append(["注意事项"])
        for tip in ("1. 请在第一个工作表「销售数据」中填写，第一行为表头，不要修改表头名称。",
                    "2. 第 2~4 行是示例数据（菜品名称含“示例”），上传前必须删除。",
                    f"3. 单个文件不超过 {MAX_IMPORT_BYTES // 1024 // 1024}MB、{MAX_IMPORT_ROWS} 行。",
                    "4. 同一份数据重复上传会被系统拦截。"):
            ws2.append([tip])
        for col, w in zip("ABCD", (14, 12, 44, 18)):
            ws2.column_dimensions[col].width = w
        for c in ws2[1]:
            c.font = Font(bold=True)
        bio = io.BytesIO()
        wb.save(bio)
        data = bio.getvalue()
        mime, fname = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "销售数据导入模板.xlsx"
    return StreamingResponse(io.BytesIO(data), media_type=mime,
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


# ---------- 6.2 上传后校验预览（不写库） ----------
@app.post("/api/sales-import/preview", tags=["销售导入"], summary="校验并预览上传文件（不写库）")
def sales_import_preview(file: UploadFile = File(...), store_id: int = Form(1), db: Session = Depends(get_db)):
    parsed = _parse_sales_file(file.filename, file.file.read())
    if not parsed["ok"]:
        return _format_error_response(parsed)
    rows = parsed["rows"]
    dish_map = _dish_sku_map(db, store_id)
    unmatched = Counter()
    unmatched_name = {}
    matched = 0
    out_rows = []
    for r in rows:
        d = dish_map.get(r["pos_sku"].upper())
        if d:
            matched += 1
        else:
            unmatched[r["pos_sku"]] += 1
            unmatched_name.setdefault(r["pos_sku"], r["dish_name"])
        if len(out_rows) < 500:
            out_rows.append({"row_no": r["row_no"], "sale_date": r["sale_date"].isoformat(), "pos_sku": r["pos_sku"],
                             "dish_name": r["dish_name"], "system_dish": d.name if d else None,
                             "quantity": float(r["quantity"]), "amount": float(r["amount"]),
                             "order_no": r["order_no"], "matched": bool(d)})
    return {
        "ok": True, "file_name": file.filename,
        "summary": {
            "total_rows": len(rows),
            "total_quantity": float(sum(r["quantity"] for r in rows)),
            "total_amount": float(sum(r["amount"] for r in rows)),
            "date_min": min(r["sale_date"] for r in rows).isoformat(),
            "date_max": max(r["sale_date"] for r in rows).isoformat(),
            "order_count": len({r["order_no"] for r in rows if r["order_no"]}),
            "matched_rows": matched, "unmatched_rows": len(rows) - matched,
            "unmatched_skus": [{"pos_sku": k, "dish_name": unmatched_name.get(k, ""), "rows": v}
                               for k, v in unmatched.most_common(20)],
        },
        "rows": out_rows, "rows_truncated": len(rows) > len(out_rows),
    }


# ---------- 6.3 确认导入：重新校验 → 防重复 → 写库 → 自动记录批次 ----------
@app.post("/api/sales-import/commit", tags=["销售导入"], summary="确认导入")
def sales_import_commit(
    file: UploadFile = File(...), store_id: int = Form(1), uploaded_by: str = Form(""),
    deduct_stock: bool = Form(True), db: Session = Depends(get_db),
):
    content = file.file.read()
    parsed = _parse_sales_file(file.filename, content)
    if not parsed["ok"]:
        return _format_error_response(parsed)
    rows = parsed["rows"]
    file_hash = hashlib.sha256(content).hexdigest()

    dup = _check_duplicate(db, store_id, rows, file_hash)
    if dup:
        if dup["level"] == "full":
            b = dup["batch"]
            extra = f"（与导入记录「{b.file_name}」内容相同）" if b else ""
            return {"ok": False, "code": "duplicate", "level": "full",
                    "message": f"该销售数据已经上传，无需再次上传{extra}", "dup_rows": dup["dup_rows"]}
        return {"ok": False, "code": "duplicate", "level": "partial", "dup_rows": dup["dup_rows"],
                "message": f"文件中有 {dup['dup_rows']} 行订单已存在于系统，为避免重复计入，请剔除这些行后重新上传"}

    try:
        dish_map = _dish_sku_map(db, store_id)
        dmin = min(r["sale_date"] for r in rows)
        dmax = max(r["sale_date"] for r in rows)
        matched = sum(1 for r in rows if dish_map.get(r["pos_sku"].upper()))
        batch = PosImportBatch(
            store_id=store_id, file_name=(file.filename or "未命名")[:255], import_date=dmin, sale_date_end=dmax,
            status="已完成", total_rows=len(rows), matched_rows=matched, unmatched_rows=len(rows) - matched,
            uploaded_by=(uploaded_by or "").strip()[:50] or None, file_hash=file_hash,
            total_amount=sum(r["amount"] for r in rows), total_quantity=sum(r["quantity"] for r in rows),
            stock_deducted=False,
        )
        db.add(batch)
        db.flush()
        sold = []
        for r in rows:
            d = dish_map.get(r["pos_sku"].upper())
            db.add(SalesRecord(store_id=store_id, batch_id=batch.id, sale_date=r["sale_date"],
                               dish_pos_sku=r["pos_sku"], dish_id=d.id if d else None,
                               quantity=r["quantity"], amount=r["amount"], order_no=r["order_no"] or None))
            sold.append((d.id if d else None, r["quantity"], r["sale_date"]))
        db.flush()
        stock_info = None
        if deduct_stock:
            stock_info = _deduct_stock(db, store_id, batch.id, sold)
            batch.stock_deducted = True
        _refresh_revenue_daily(db, store_id, {r["sale_date"] for r in rows})
        db.commit()
    except Exception as e:
        db.rollback()
        logger.exception("销售数据导入失败")
        raise HTTPException(status_code=500, detail=f"写入数据库失败，已整体回滚：{e}")
    return {"ok": True, "code": "success", "message": "数据上传成功", "batch_id": batch.id,
            "total_rows": len(rows), "matched_rows": matched, "unmatched_rows": len(rows) - matched,
            "stock": stock_info}


# ---------- 6.4 批次详情 / 撤销 ----------
@app.get("/api/pos-import-batches/{item_id}/detail", tags=["POS 导入批次"], summary="导入批次详情")
def import_batch_detail(item_id: int, db: Session = Depends(get_db)):
    b = db.query(PosImportBatch).filter(PosImportBatch.id == item_id).first()
    if not b:
        raise HTTPException(status_code=404, detail="导入记录不存在")
    recs = db.query(SalesRecord).filter(SalesRecord.batch_id == b.id).order_by(SalesRecord.id).limit(1000).all()
    names = {d.id: d.name for d in db.query(Dish).filter(Dish.store_id == b.store_id).all()}
    unmatched = Counter(r.dish_pos_sku for r in db.query(SalesRecord).filter(
        SalesRecord.batch_id == b.id, SalesRecord.dish_id.is_(None)).all())
    so = db.query(StockOut).filter(StockOut.store_id == b.store_id, StockOut.out_type == "销售自动扣减",
                                   StockOut.source_id == b.id).first()
    stock = None
    if so:
        items = db.query(StockOutItem).filter(StockOutItem.stock_out_id == so.id).all()
        stock = {"stock_out_id": so.id, "materials": len(items),
                 "cost_amount": round(sum(float(i.cost_amount or 0) for i in items), 2)}
    return {
        "batch": {"id": b.id, "file_name": b.file_name, "import_date": b.import_date, "sale_date_end": b.sale_date_end,
                  "total_rows": b.total_rows, "matched_rows": b.matched_rows, "unmatched_rows": b.unmatched_rows,
                  "uploaded_by": b.uploaded_by, "total_amount": float(b.total_amount or 0),
                  "total_quantity": float(b.total_quantity or 0), "stock_deducted": bool(b.stock_deducted),
                  "created_at": b.created_at, "status": b.status},
        "rows": [{"sale_date": r.sale_date, "pos_sku": r.dish_pos_sku, "dish": names.get(r.dish_id),
                  "quantity": float(r.quantity or 0), "amount": float(r.amount or 0), "order_no": r.order_no}
                 for r in recs],
        "rows_truncated": (b.total_rows or 0) > len(recs),
        "unmatched_skus": [{"pos_sku": k, "rows": v} for k, v in unmatched.most_common()],
        "stock": stock,
    }


@app.delete("/api/pos-import-batches/{item_id}", tags=["POS 导入批次"], summary="撤销导入批次（回滚销售流水与库存）")
def delete_import_batch(item_id: int, db: Session = Depends(get_db)):
    b = db.query(PosImportBatch).filter(PosImportBatch.id == item_id).first()
    if not b:
        raise HTTPException(status_code=404, detail="导入记录不存在")
    try:
        _rollback_batch(db, b)
        db.commit()
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"撤销失败：{e}")
    return {"ok": True, "deleted_id": item_id}


# ============================================================
#  7. 库存盘点
# ============================================================
@app.post("/api/stock-take/submit", tags=["盘点单"], summary="提交盘点：生成盘点单并按实盘数校正库存")
def submit_stock_take(payload: StockTakeSubmitIn, db: Session = Depends(get_db)):
    if not payload.items:
        raise HTTPException(status_code=400, detail="没有需要提交的盘点明细")
    mats = {m.id: m for m in db.query(Material).filter(Material.store_id == payload.store_id).all()}
    take = StockTake(store_id=payload.store_id, warehouse_id=payload.warehouse_id, take_date=payload.take_date,
                     status="approved", operator_id=payload.operator_id)
    db.add(take)
    db.flush()
    changed, diff_amount = 0, 0.0
    for line in payload.items:
        if line.actual_quantity < 0:
            db.rollback()
            raise HTTPException(status_code=400, detail="实盘数量不能为负数")
        inv = db.query(Inventory).filter(
            Inventory.store_id == payload.store_id, Inventory.warehouse_id == payload.warehouse_id,
            Inventory.material_id == line.material_id).first()
        book = float(inv.quantity or 0) if inv else 0.0
        diff = line.actual_quantity - book
        m = mats.get(line.material_id)
        unit_cost = float(inv.avg_cost or 0) if inv else float(m.latest_purchase_price or 0) if m else 0.0
        db.add(StockTakeItem(stock_take_id=take.id, material_id=line.material_id, book_quantity=book,
                             actual_quantity=line.actual_quantity, diff_quantity=diff))
        if inv:
            inv.quantity = line.actual_quantity
        else:
            db.add(Inventory(store_id=payload.store_id, warehouse_id=payload.warehouse_id,
                             material_id=line.material_id, quantity=line.actual_quantity, avg_cost=unit_cost))
        if abs(diff) > 1e-9:
            changed += 1
        diff_amount += diff * unit_cost
    db.commit()
    return {"ok": True, "id": take.id, "items": len(payload.items), "changed": changed,
            "diff_amount": round(diff_amount, 2)}


# ============================================================
#  8. 报表中心（统一接口 GET /api/reports/{key}?start=&end=）
#  返回结构：{title, columns[{key,label,type}], rows[], summary[{label,value,type,tone}], chart, note}
#  type ∈ text | int | qty | money | pct | date
# ============================================================
def _col(key, label, typ="text"):
    return {"key": key, "label": label, "type": typ}


def _sm(label, value, typ="money", tone=None):
    return {"label": label, "value": value, "type": typ, "tone": tone}


def _period(start: Optional[str], end: Optional[str]):
    try:
        e = date.fromisoformat(end) if end else date.today()
        s = date.fromisoformat(start) if start else e - timedelta(days=29)
    except ValueError:
        raise HTTPException(status_code=400, detail="日期格式应为 YYYY-MM-DD")
    if s > e:
        raise HTTPException(status_code=400, detail="开始日期不能晚于结束日期")
    return s, e


def _safe_div(a, b):
    return a / b if b else 0.0


def _sales_by_dish(db, store_id, s, e):
    rows = db.query(SalesRecord.dish_id, SalesRecord.dish_pos_sku,
                    func.sum(SalesRecord.quantity), func.sum(SalesRecord.amount)).filter(
        SalesRecord.store_id == store_id, SalesRecord.sale_date >= s, SalesRecord.sale_date <= e
    ).group_by(SalesRecord.dish_id, SalesRecord.dish_pos_sku).all()
    return [(d, sku, float(q or 0), float(a or 0)) for d, sku, q, a in rows]


def _dish_unit_costs(db, store_id):
    bom = _dish_bom(db, store_id)
    _, cost = _material_cost_map(db, store_id)
    return {did: sum(q * (1 + loss) * cost.get(mid, 0.0) for mid, q, loss in lines) for did, lines in bom.items()}


def _moves(db, store_id, kind, lo, hi):
    """出入库汇总 {material_id: 数量}。入库只统计已审批单据。"""
    if kind == "in":
        mcol = StockInItem.material_id
        q = db.query(mcol, func.sum(StockInItem.quantity)).join(StockIn, StockIn.id == StockInItem.stock_in_id) \
            .filter(StockIn.store_id == store_id, StockIn.status == "approved")
        ts = StockIn.created_at
    else:
        mcol = StockOutItem.material_id
        q = db.query(mcol, func.sum(StockOutItem.quantity)).join(StockOut, StockOut.id == StockOutItem.stock_out_id) \
            .filter(StockOut.store_id == store_id)
        ts = StockOut.created_at
    if lo is not None:
        q = q.filter(ts >= lo)
    if hi is not None:
        q = q.filter(ts < hi)
    return {mid: float(v or 0) for mid, v in q.group_by(mcol).all()}


def _collect_alerts(db, store_id):
    mats, cost = _material_cost_map(db, store_id)
    inv = defaultdict(float)
    for i in db.query(Inventory).filter(Inventory.store_id == store_id).all():
        inv[i.material_id] += float(i.quantity or 0)
    alerts = []
    for mid, qty in inv.items():
        m = mats.get(mid)
        if not m:
            continue
        safety = float(m.safety_stock or 0)
        if qty < 0:
            alerts.append({"type": "负库存", "level": "danger", "material": m.name, "unit": m.unit, "qty": qty,
                           "detail": f"账面 {qty:g}，可能是采购未入库或配方用量偏大"})
        elif safety > 0 and qty <= safety:
            alerts.append({"type": "低库存", "level": "warn", "material": m.name, "unit": m.unit, "qty": qty,
                           "detail": f"当前 {qty:g}，安全库存 {safety:g}"})
    horizon = date.today() + timedelta(days=7)
    batches = db.query(StockInItem, StockIn).select_from(StockInItem).join(StockIn, StockIn.id == StockInItem.stock_in_id).filter(
        StockIn.store_id == store_id, StockIn.status == "approved",
        StockInItem.expire_date.is_not(None), StockInItem.expire_date <= horizon).order_by(StockInItem.expire_date).all()
    for it, si in batches:
        m = mats.get(it.material_id)
        if not m or inv.get(it.material_id, 0) <= 0:
            continue                                   # 该物料已无库存，批次提醒没有意义
        left = (it.expire_date - date.today()).days
        alerts.append({"type": "已过期" if left < 0 else "临期", "level": "danger" if left < 0 else "warn",
                       "material": m.name, "unit": m.unit, "qty": float(it.quantity or 0),
                       "detail": f"入库 {float(it.quantity or 0):g}{m.unit}，" +
                                 (f"已过期 {-left} 天" if left < 0 else f"{left} 天后到期") + f"（{it.expire_date}）"})
    return alerts


def report_daily(db, store_id, s, e):
    mats, cost = _material_cost_map(db, store_id)
    inv = defaultdict(float)
    for i in db.query(Inventory).filter(Inventory.store_id == store_id).all():
        inv[i.material_id] += float(i.quantity or 0)
    lo = datetime.combine(s, datetime.min.time())
    hi = datetime.combine(e + timedelta(days=1), datetime.min.time())
    in_r, out_r = _moves(db, store_id, "in", lo, hi), _moves(db, store_id, "out", lo, hi)
    in_a, out_a = _moves(db, store_id, "in", hi, None), _moves(db, store_id, "out", hi, None)
    rows = []
    tot_in = tot_out = tot_close = 0.0
    for mid in set(inv) | set(in_r) | set(out_r):
        m = mats.get(mid)
        if not m:
            continue
        closing = inv.get(mid, 0.0) - in_a.get(mid, 0.0) + out_a.get(mid, 0.0)
        opening = closing - in_r.get(mid, 0.0) + out_r.get(mid, 0.0)
        c = cost.get(mid, 0.0)
        tot_in += in_r.get(mid, 0.0) * c
        tot_out += out_r.get(mid, 0.0) * c
        tot_close += closing * c
        rows.append({"material": m.name, "unit": m.unit, "opening": opening, "in": in_r.get(mid, 0.0),
                     "out": out_r.get(mid, 0.0), "closing": closing, "closing_amount": closing * c})
    rows.sort(key=lambda r: -r["closing_amount"])
    return {
        "title": "进销存日报",
        "columns": [_col("material", "原材料"), _col("unit", "单位"), _col("opening", "期初", "qty"),
                    _col("in", "入库", "qty"), _col("out", "出库", "qty"), _col("closing", "期末", "qty"),
                    _col("closing_amount", "期末金额", "money")],
        "rows": rows,
        "summary": [_sm("期间入库金额", tot_in), _sm("期间出库金额", tot_out), _sm("期末库存金额", tot_close, tone="green"),
                    _sm("物料品项", len(rows), "int")],
        "chart": {"type": "hbar", "title": "期末库存金额 Top10",
                  "labels": [r["material"] for r in rows[:10]], "series": [{"name": "期末金额", "values": [round(r["closing_amount"], 2) for r in rows[:10]]}]},
        "note": "期初/期末由当前库存按期间内出入库倒推；仅统计已审批入库单，盘点校正不计入出入库。",
    }


def report_loss(db, store_id, s, e):
    mats, cost = _material_cost_map(db, store_id)
    bom = _dish_bom(db, store_id)
    theo, recipe_loss = defaultdict(float), defaultdict(float)
    for did, _, qty, _ in _sales_by_dish(db, store_id, s, e):
        for mid, q, loss in bom.get(did, []):
            theo[mid] += qty * q
            recipe_loss[mid] += qty * q * loss
    diffs = defaultdict(float)
    q = db.query(StockTakeItem.material_id, func.sum(StockTakeItem.diff_quantity)).join(
        StockTake, StockTake.id == StockTakeItem.stock_take_id).filter(
        StockTake.store_id == store_id, StockTake.take_date >= s, StockTake.take_date <= e
    ).group_by(StockTakeItem.material_id)
    for mid, v in q.all():
        diffs[mid] = float(v or 0)
    rows, loss_amt, over_amt = [], 0.0, 0.0
    for mid in set(theo) | set(diffs):
        m = mats.get(mid)
        if not m:
            continue
        d_amt = diffs.get(mid, 0.0) * cost.get(mid, 0.0)
        if d_amt < 0:
            loss_amt += -d_amt
        else:
            over_amt += d_amt
        t = theo.get(mid, 0.0)
        rows.append({"material": m.name, "unit": m.unit, "theory": t, "recipe_loss": recipe_loss.get(mid, 0.0),
                     "diff_qty": diffs.get(mid, 0.0), "diff_amount": d_amt,
                     "diff_rate": _safe_div(-diffs.get(mid, 0.0), t + recipe_loss.get(mid, 0.0))})
    rows.sort(key=lambda r: r["diff_amount"])
    return {
        "title": "损耗分析",
        "columns": [_col("material", "原材料"), _col("unit", "单位"), _col("theory", "理论用量", "qty"),
                    _col("recipe_loss", "配方损耗", "qty"), _col("diff_qty", "盘点差异", "qty"),
                    _col("diff_amount", "差异金额", "money"), _col("diff_rate", "异常损耗率", "pct")],
        "rows": rows,
        "summary": [_sm("盘亏金额", loss_amt, tone="red"), _sm("盘盈金额", over_amt, tone="green"),
                    _sm("涉及物料", len(rows), "int")],
        "chart": {"type": "hbar", "title": "盘点差异金额（负数为盘亏）",
                  "labels": [r["material"] for r in rows[:10]], "series": [{"name": "差异金额", "values": [round(r["diff_amount"], 2) for r in rows[:10]]}]},
        "note": "理论用量 = 销量 × 配方单份用量；配方损耗按配方损耗率折算；盘点差异 = 实盘 − 账面（负数即盘亏）。请先在「库存盘点」录入数据。",
    }


def report_gross(db, store_id, s, e):
    names = {d.id: d.name for d in db.query(Dish).filter(Dish.store_id == store_id).all()}
    ucost = _dish_unit_costs(db, store_id)
    rows, t_rev, t_cost = [], 0.0, 0.0
    for did, sku, qty, amt in _sales_by_dish(db, store_id, s, e):
        uc = ucost.get(did, 0.0)
        c = uc * qty
        t_rev += amt
        t_cost += c
        name = names.get(did) or f"未匹配（{sku}）"
        rows.append({"dish": name, "qty": qty, "revenue": amt, "unit_cost": uc, "cost": c, "profit": amt - c,
                     "margin": _safe_div(amt - c, amt), "flag": "" if did in ucost else "无配方"})
    rows.sort(key=lambda r: -r["profit"])
    return {
        "title": "毛利分析",
        "columns": [_col("dish", "菜品"), _col("qty", "销量", "qty"), _col("revenue", "销售额", "money"),
                    _col("unit_cost", "单份成本", "money"), _col("cost", "总成本", "money"),
                    _col("profit", "毛利", "money"), _col("margin", "毛利率", "pct"), _col("flag", "提示")],
        "rows": rows,
        "summary": [_sm("销售额", t_rev), _sm("理论原料成本", t_cost), _sm("毛利", t_rev - t_cost, tone="green"),
                    _sm("综合毛利率", _safe_div(t_rev - t_cost, t_rev), "pct", "green")],
        "chart": {"type": "hbar", "title": "毛利 Top10", "labels": [r["dish"] for r in rows[:10]],
                  "series": [{"name": "毛利", "values": [round(r["profit"], 2) for r in rows[:10]]}]},
        "note": "成本 = 配方用量 ×（1+损耗率）× 物料当前加权平均成本；标注“无配方”的菜品成本按 0 计算，请到「菜品与配方」补录。",
    }


def report_supplier(db, store_id, s, e):
    lo = datetime.combine(s, datetime.min.time())
    hi = datetime.combine(e + timedelta(days=1), datetime.min.time())
    q = db.query(PurchaseOrder.supplier_id, func.sum(PurchaseOrderItem.amount),
                 func.count(func.distinct(PurchaseOrder.id))).select_from(PurchaseOrder).join(
        PurchaseOrderItem, PurchaseOrderItem.purchase_order_id == PurchaseOrder.id).filter(
        PurchaseOrder.store_id == store_id, PurchaseOrder.status == "approved",
        PurchaseOrder.created_at >= lo, PurchaseOrder.created_at < hi).group_by(PurchaseOrder.supplier_id)
    names = {x.id: x.name for x in db.query(Supplier).filter(Supplier.store_id == store_id).all()}
    data = [(names.get(sid, f"供应商#{sid}"), float(a or 0), int(c)) for sid, a, c in q.all()]
    data.sort(key=lambda x: -x[1])
    total = sum(x[1] for x in data)
    rows = [{"rank": i + 1, "supplier": n, "orders": c, "amount": a, "share": _safe_div(a, total),
             "avg": _safe_div(a, c)} for i, (n, a, c) in enumerate(data)]
    return {
        "title": "供应商采购排行",
        "columns": [_col("rank", "排名", "int"), _col("supplier", "供应商"), _col("orders", "采购单数", "int"),
                    _col("amount", "采购金额", "money"), _col("share", "占比", "pct"), _col("avg", "单均金额", "money")],
        "rows": rows,
        "summary": [_sm("采购总额", total), _sm("供应商数", len(rows), "int"), _sm("采购单数", sum(r["orders"] for r in rows), "int")],
        "chart": {"type": "hbar", "title": "采购金额 Top10", "labels": [r["supplier"] for r in rows[:10]],
                  "series": [{"name": "采购金额", "values": [round(r["amount"], 2) for r in rows[:10]]}]},
        "note": "仅统计已审批的采购单。",
    }


def report_abc(db, store_id, s, e):
    mats, cost = _material_cost_map(db, store_id)
    lo = datetime.combine(s, datetime.min.time())
    hi = datetime.combine(e + timedelta(days=1), datetime.min.time())
    q = db.query(StockOutItem.material_id, func.sum(StockOutItem.cost_amount)).join(
        StockOut, StockOut.id == StockOutItem.stock_out_id).filter(
        StockOut.store_id == store_id, StockOut.created_at >= lo, StockOut.created_at < hi
    ).group_by(StockOutItem.material_id)
    val = {mid: float(v or 0) for mid, v in q.all()}
    basis = "期间出库成本"
    if not any(val.values()):
        basis = "当前库存金额"
        val = defaultdict(float)
        for i in db.query(Inventory).filter(Inventory.store_id == store_id).all():
            val[i.material_id] += max(float(i.quantity or 0), 0) * cost.get(i.material_id, 0.0)
    data = sorted(((mats[mid].name, mats[mid].unit, v) for mid, v in val.items() if mid in mats and v > 0),
                  key=lambda x: -x[2])
    total, cum, rows = sum(x[2] for x in data), 0.0, []
    counts = {"A": 0, "B": 0, "C": 0}
    for name, unit, v in data:
        cum += v
        share_cum = _safe_div(cum, total)
        cls = "A" if share_cum - _safe_div(v, total) < 0.7 else "B" if share_cum - _safe_div(v, total) < 0.9 else "C"
        counts[cls] += 1
        rows.append({"cls": cls, "material": name, "unit": unit, "amount": v, "share": _safe_div(v, total), "cum": share_cum})
    return {
        "title": "ABC 分析",
        "columns": [_col("cls", "类别"), _col("material", "原材料"), _col("unit", "单位"), _col("amount", "金额", "money"),
                    _col("share", "占比", "pct"), _col("cum", "累计占比", "pct")],
        "rows": rows,
        "summary": [_sm("统计口径", basis, "text"), _sm("A 类（累计≤70%）", counts["A"], "int", "red"),
                    _sm("B 类（70%~90%）", counts["B"], "int"), _sm("C 类（其余）", counts["C"], "int")],
        "chart": {"type": "hbar", "title": f"{basis} Top10", "labels": [r["material"] for r in rows[:10]],
                  "series": [{"name": "金额", "values": [round(r["amount"], 2) for r in rows[:10]]}]},
        "note": "A 类物料是成本大头，应重点控制采购价与损耗；C 类可放宽管理。",
    }


def _daily_series(db, store_id, s, e):
    """按天汇总销售额/订单/销量，返回 {date: {...}}"""
    q = db.query(SalesRecord.sale_date, func.sum(SalesRecord.amount), func.sum(SalesRecord.quantity),
                 func.count(SalesRecord.id), func.count(func.distinct(SalesRecord.order_no))).filter(
        SalesRecord.store_id == store_id, SalesRecord.sale_date >= s, SalesRecord.sale_date <= e
    ).group_by(SalesRecord.sale_date)
    out = {}
    for d, amt, qty, cnt, orders in q.all():
        out[d] = {"revenue": float(amt or 0), "qty": float(qty or 0), "orders": int(orders) if orders else int(cnt)}
    return out


def report_sales(db, store_id, s, e):
    series = _daily_series(db, store_id, s, e)
    days = sorted(series)
    rows = [{"date": d.isoformat(), "revenue": series[d]["revenue"], "orders": series[d]["orders"],
             "qty": series[d]["qty"], "avg": _safe_div(series[d]["revenue"], series[d]["orders"])} for d in days]
    total = sum(r["revenue"] for r in rows)
    orders = sum(r["orders"] for r in rows)
    best = max(rows, key=lambda r: r["revenue"]) if rows else None
    return {
        "title": "销售趋势",
        "columns": [_col("date", "日期", "date"), _col("revenue", "营业额", "money"), _col("orders", "订单数", "int"),
                    _col("qty", "销量", "qty"), _col("avg", "客单价", "money")],
        "rows": rows,
        "summary": [_sm("营业额", total, tone="green"), _sm("订单数", orders, "int"), _sm("客单价", _safe_div(total, orders)),
                    _sm("日均营业额", _safe_div(total, len(rows))),
                    _sm("最高单日", f"{best['date']}  ¥{best['revenue']:,.0f}" if best else "—", "text")],
        "chart": {"type": "line", "title": "每日营业额", "labels": [r["date"][5:] for r in rows],
                  "series": [{"name": "营业额", "values": [round(r["revenue"], 2) for r in rows]}]},
        "note": "数据来自「销售数据导入」；无订单号的文件按行数估算订单数。",
    }


def report_pnl(db, store_id, s, e):
    series = _daily_series(db, store_id, s, e)
    costs = {c.cost_date: c for c in db.query(CostDaily).filter(
        CostDaily.store_id == store_id, CostDaily.cost_date >= s, CostDaily.cost_date <= e).all()}
    ucost = _dish_unit_costs(db, store_id)
    theo = defaultdict(float)
    for d, did, qty in db.query(SalesRecord.sale_date, SalesRecord.dish_id, func.sum(SalesRecord.quantity)).filter(
            SalesRecord.store_id == store_id, SalesRecord.sale_date >= s, SalesRecord.sale_date <= e
    ).group_by(SalesRecord.sale_date, SalesRecord.dish_id).all():
        theo[d] += ucost.get(did, 0.0) * float(qty or 0)
    rows = []
    for d in sorted(set(series) | set(costs)):
        rev = series.get(d, {}).get("revenue", 0.0)
        c = costs.get(d)
        mat = float(c.material_cost or 0) if c and float(c.material_cost or 0) > 0 else theo.get(d, 0.0)
        opex = sum(float(getattr(c, k) or 0) for k in ("labor_cost", "rent_cost", "utility_cost", "depreciation_cost", "other_cost")) if c else 0.0
        rows.append({"date": d.isoformat(), "revenue": rev, "material": mat, "gross": rev - mat,
                     "gross_rate": _safe_div(rev - mat, rev), "opex": opex, "profit": rev - mat - opex})
    t = lambda k: sum(r[k] for r in rows)
    return {
        "title": "简化损益表",
        "columns": [_col("date", "日期", "date"), _col("revenue", "营业收入", "money"), _col("material", "原料成本", "money"),
                    _col("gross", "毛利", "money"), _col("gross_rate", "毛利率", "pct"), _col("opex", "人工房租水电等", "money"),
                    _col("profit", "税前利润", "money")],
        "rows": rows,
        "summary": [_sm("营业收入", t("revenue")), _sm("原料成本", t("material")), _sm("其他费用", t("opex")),
                    _sm("税前利润", t("profit"), tone="green" if t("profit") >= 0 else "red"),
                    _sm("净利率", _safe_div(t("profit"), t("revenue")), "pct")],
        "chart": {"type": "line", "title": "收入 / 成本 / 利润趋势", "labels": [r["date"][5:] for r in rows],
                  "series": [{"name": "收入", "values": [round(r["revenue"], 2) for r in rows]},
                             {"name": "原料成本", "values": [round(r["material"], 2) for r in rows]},
                             {"name": "利润", "values": [round(r["profit"], 2) for r in rows]}]},
        "note": "原料成本优先取「财务管理」当日录入值，未录入则用销量×配方理论成本；费用取当日录入的人工、房租、水电、折旧、其他。",
    }


def report_alert(db, store_id, s, e):
    alerts = _collect_alerts(db, store_id)
    order = {"danger": 0, "warn": 1}
    alerts.sort(key=lambda a: order.get(a["level"], 2))
    rows = [{"type": a["type"], "material": a["material"], "detail": a["detail"]} for a in alerts]
    cnt = Counter(a["type"] for a in alerts)
    return {
        "title": "库存预警",
        "columns": [_col("type", "类型"), _col("material", "原材料"), _col("detail", "详情")],
        "rows": rows,
        "summary": [_sm("低库存", cnt["低库存"], "int", "red" if cnt["低库存"] else None),
                    _sm("负库存", cnt["负库存"], "int", "red" if cnt["负库存"] else None),
                    _sm("临期（7天内）", cnt["临期"], "int"), _sm("已过期", cnt["已过期"], "int", "red" if cnt["已过期"] else None)],
        "chart": None,
        "note": "低库存依据原材料的「安全库存」；临期依据入库时填写的保质期，且仅提示仍有库存的物料。本报表不受日期范围影响。",
    }


REPORTS = {"daily": report_daily, "loss": report_loss, "gross": report_gross, "supplier": report_supplier,
           "abc": report_abc, "sales": report_sales, "pnl": report_pnl, "alert": report_alert}


@app.get("/api/reports/{key}", tags=["报表"], summary="报表数据")
def get_report(key: str, start: Optional[str] = None, end: Optional[str] = None, store_id: int = 1,
               db: Session = Depends(get_db)):
    fn = REPORTS.get(key)
    if not fn:
        raise HTTPException(status_code=404, detail="报表不存在")
    s, e = _period(start, end)
    out = fn(db, store_id, s, e)
    out.update({"key": key, "start": s.isoformat(), "end": e.isoformat()})
    return out


# ============================================================
#  9. 首页看板
# ============================================================
@app.get("/api/dashboard/summary", tags=["系统"], summary="首页看板汇总")
def dashboard_summary(store_id: int = 1, db: Session = Depends(get_db)):
    today = date.today()
    series = _daily_series(db, store_id, today - timedelta(days=6), today)
    trend = [{"date": (today - timedelta(days=i)).isoformat(),
              "revenue": series.get(today - timedelta(days=i), {}).get("revenue", 0.0)} for i in range(6, -1, -1)]
    t, y = series.get(today, {}), series.get(today - timedelta(days=1), {})
    cnt = lambda m: db.query(func.count(m.id)).filter(m.store_id == store_id).scalar() or 0
    alerts = _collect_alerts(db, store_id)
    recent = db.query(PosImportBatch).filter(PosImportBatch.store_id == store_id).order_by(PosImportBatch.id.desc()).limit(5).all()
    last_sale = db.query(func.max(SalesRecord.sale_date)).filter(SalesRecord.store_id == store_id).scalar()
    return {
        "today": {"revenue": t.get("revenue", 0.0), "orders": t.get("orders", 0)},
        "yesterday": {"revenue": y.get("revenue", 0.0), "orders": y.get("orders", 0)},
        "trend": trend,
        "counts": {"suppliers": cnt(Supplier), "purchase_orders": cnt(PurchaseOrder), "dishes": cnt(Dish), "inventory": cnt(Inventory)},
        "pending": {
            "purchase_orders": db.query(func.count(PurchaseOrder.id)).filter(PurchaseOrder.store_id == store_id, PurchaseOrder.status == "pending").scalar() or 0,
            "stock_in": db.query(func.count(StockIn.id)).filter(StockIn.store_id == store_id, StockIn.status == "pending").scalar() or 0,
        },
        "alerts": alerts[:8], "alert_total": len(alerts),
        "last_sale_date": last_sale,
        "recent_imports": [{"id": b.id, "file_name": b.file_name, "import_date": b.import_date,
                            "sale_date_end": b.sale_date_end, "total_rows": b.total_rows,
                            "uploaded_by": b.uploaded_by} for b in recent],
    }


# ============================================================
#  系统接口
# ============================================================
@app.get("/", include_in_schema=False)
def root():
    return FileResponse("static/index.html")


@app.get("/api/health", tags=["系统"])
def health(db: Session = Depends(get_db)):
    try:
        db.execute(text("SELECT 1"))
        return {"status": "ok", "database": "connected"}
    except Exception as e:
        return {"status": "error", "detail": str(e)}


if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", 8000))
    uvicorn.run("main:app", host="0.0.0.0", port=port, reload=False)
